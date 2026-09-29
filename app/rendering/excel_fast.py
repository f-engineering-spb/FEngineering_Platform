"""Fast Excel preview: parse the workbook ONCE, cache every sheet.

Sheet switching on the frontend is served from the in-memory LRU without
re-reading the file. Modern .xlsx/.xlsm go through openpyxl (read-only),
legacy .xls through xlrd.
"""

from __future__ import annotations

import html
import os
import threading
from functools import lru_cache

_MAX_CACHED_WORKBOOKS = 8
_MAX_ROWS_PER_SHEET = 2000
_MAX_COLS_PER_SHEET = 100


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _sheet_to_html(name, rows) -> str:
    out = ["<table><caption>%s</caption><tbody>" % html.escape(str(name))]
    for row in rows:
        out.append("<tr>%s</tr>" % "".join(
            "<td>%s</td>" % html.escape(_cell_text(c)) for c in row))
    out.append("</tbody></table>")
    return "".join(out)


def _parse_xlsx(path: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {}
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row[:_MAX_COLS_PER_SHEET]))
                if len(rows) >= _MAX_ROWS_PER_SHEET:
                    break
            sheets[str(ws.title)] = _sheet_to_html(ws.title, rows)
        return sheets
    finally:
        wb.close()


def _parse_xls(path: str) -> dict:
    import xlrd

    book = xlrd.open_workbook(path, on_demand=False)
    sheets = {}
    for idx in range(book.nsheets):
        sh = book.sheet_by_index(idx)
        rows = []
        for r in range(min(sh.nrows, _MAX_ROWS_PER_SHEET)):
            rows.append([sh.cell_value(r, c)
                         for c in range(min(sh.ncols, _MAX_COLS_PER_SHEET))])
        sheets[str(sh.name)] = _sheet_to_html(sh.name, rows)
    return sheets


class CachedWorkbook:
    """All sheets of one workbook, parsed exactly once."""

    def __init__(self, path: str, sheets: dict):
        self.path = path
        self._sheets = sheets
        self._lock = threading.Lock()

    def sheet_names(self):
        return list(self._sheets.keys())

    def sheet_html(self, name: str) -> str:
        with self._lock:
            if name in self._sheets:
                return self._sheets[name]
            first = next(iter(self._sheets), "")
            return self._sheets.get(first, "<p>empty workbook</p>")


@lru_cache(maxsize=_MAX_CACHED_WORKBOOKS)
def _load_cached(path: str, mtime_ns: int) -> CachedWorkbook:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        sheets = _parse_xls(path)
    else:
        sheets = _parse_xlsx(path)
    return CachedWorkbook(path, sheets)


def open_workbook(path: str) -> CachedWorkbook:
    """Parse once (or hit LRU) and return every cached sheet."""
    st = os.stat(str(path))
    return _load_cached(str(path), st.st_mtime_ns)
